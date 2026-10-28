import pytest
import io
import openpyxl
from app.db.models.user import User

def test_hierarchy_consolidation_endpoints(client, db, superuser_token_headers):
    """
    Test IFRS 10 group consolidation endpoints and Excel generation.
    """
    # 1. Test GET /api/v1/hierarchy/consolidation
    resp = client.get("/api/v1/hierarchy/consolidation", headers=superuser_token_headers)
    assert resp.status_code == 200
    data = resp.json()
    assert data["success"] is True
    assert data["standard"] == "IFRS 10 Consolidated Financial Statements"
    assert data["total_assets"] == 248200000
    assert data["total_liabilities"] == 96300000
    assert data["total_equity"] == 151900000
    assert data["is_balanced"] is True
    assert data["delta"] == 0.0

    # 2. Test GET /api/v1/hierarchy/consolidation/export-excel
    resp_excel = client.get("/api/v1/hierarchy/consolidation/export-excel", headers=superuser_token_headers)
    assert resp_excel.status_code == 200
    assert resp_excel.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    
    # Load and verify Excel sheets with openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(resp_excel.content))
    assert "Консолидированный_Баланс" in wb.sheetnames
    assert "Внутригрупповые_Элиминации" in wb.sheetnames
    assert "Структура_и_NCI" in wb.sheetnames
